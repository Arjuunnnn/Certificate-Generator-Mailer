import os
import re
import time
import logging
import secrets
import string
import pandas as pd
import qrcode
import comtypes.client
import smtplib
from pptx import Presentation
from pptx.util import Inches
from tqdm import tqdm
from email.message import EmailMessage
from concurrent.futures import ThreadPoolExecutor, as_completed
from pdf2image import convert_from_path       # pip install pdf2image
                                              # Also install poppler (Windows):
                                              # https://github.com/oschwartz10612/poppler-windows/releases
                                              # Add poppler/Library/bin to PATH
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request

# ---------------------------------------------
#  SETTINGS
# ---------------------------------------------
DATA_FILE       = "data.xlsx"
OUTPUT_FOLDER   = "certificates"
NAME_COLUMN     = "Name"
CERT_PREFIX     = "IEDC/2026/"
VALIDATION_URL  = "https://yourcollege.edu/verify?id="
QR_PLACEHOLDER  = "QR_PLACEHOLDER"

# Each entry: (Excel sheet name, template file, output subfolder)
CATEGORIES = [
    ("Winners",      "template_winner.pptx",      "winners"),
    ("Participants", "template_participant.pptx",  "participants"),
    ("Coordinators", "template_coordinator.pptx",  "coordinators"),
]

# ---------------------------------------------
#  GOOGLE DRIVE SETTINGS
# ---------------------------------------------
# Uses OAuth2 credentials (client_id / client_secret / refresh_token).
# Get these from Google Cloud Console:
#   1. console.cloud.google.com -> project -> Enable Drive API
#   2. APIs & Services -> Credentials -> OAuth 2.0 Client IDs
#      (Application type: Desktop app) -> download JSON
#   3. Run the OAuth consent flow once to obtain a refresh token,
#      then paste the three values below.
#
# Folder structure created automatically each run:
#   GDRIVE_FOLDER_ID/
#     └── <event_title>/          <- entered at runtime
#           ├── winners/
#           ├── participants/
#           └── coordinators/
GDRIVE_CLIENT_ID     = "197710549184-aaf7oi2heuu6b7c82bqjf2u9h34kviuu.apps.googleusercontent.com"      # 🔴 Replace
GDRIVE_CLIENT_SECRET = "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"  # 🔴 Replace
GDRIVE_REFRESH_TOKEN = "1//0gn7Ae2eVSasMCgYIARAAGBASNwF-L9IrZ_Yr3NuiOkSsEmDBiSUu5rfuzkEcut0-2H202U2yndw-PvyFCXRNwU0ivIbmnDx5JIc"  # 🔴 Replace
GDRIVE_FOLDER_ID     = "1noDresKVDkhKxoBicqz-O_ikQmR8ob_B"  # 🔴 From Drive folder URL
GDRIVE_TOKEN_URI     = "https://oauth2.googleapis.com/token"
GDRIVE_SCOPES        = ["https://www.googleapis.com/auth/drive"]
 
# EMAIL SETTINGS
SEND_EMAILS         = True
EMAIL_USER          = "your_email@zohomail.in"
EMAIL_PASS          = "your_app_password_here"
EMAIL_TEMPLATE_FILE = "email.txt"
ZOHO_SMTP_HOST      = "smtp.zoho.in"
ZOHO_SMTP_PORT      = 465
 
# OPTIONS
PNG_DPI             = 300
MAX_EMAIL_RETRIES   = 3
RETRY_DELAY_SECONDS = 5
EMAIL_THREADS       = 3
 
# ---------------------------------------------
#  LOGGING
# ---------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler()]
)
log = logging.getLogger(__name__)
 
# ---------------------------------------------
#  UNIQUE ID
# ---------------------------------------------
def generate_unique_suffix(length=4):
    alphabet = string.ascii_uppercase + string.digits
    return "".join(secrets.choice(alphabet) for _ in range(length))
 
# ---------------------------------------------
#  GOOGLE DRIVE HELPERS
# ---------------------------------------------
def get_drive_service():
    """Build an authenticated Drive client using OAuth2 refresh token credentials.
 
    NOTE: The refresh token must have been generated with the same scope as
    GDRIVE_SCOPES. Do NOT pass scopes here — let the token carry its own scope,
    which avoids the invalid_scope error when the token was issued without an
    explicit scope list in the original OAuth consent flow.
    """
    creds = Credentials(
        token=None,
        refresh_token=GDRIVE_REFRESH_TOKEN,
        token_uri=GDRIVE_TOKEN_URI,
        client_id=GDRIVE_CLIENT_ID,
        client_secret=GDRIVE_CLIENT_SECRET,
        # scopes intentionally omitted — use whatever scope the token was issued with
    )
    creds.refresh(Request())
    return build("drive", "v3", credentials=creds)
 
 
def get_or_create_drive_folder(drive_service, name, parent_id):
    q = (
        f"name='{name}' and '{parent_id}' in parents "
        f"and mimeType='application/vnd.google-apps.folder' and trashed=false"
    )
    results = drive_service.files().list(q=q, fields="files(id)").execute()
    files   = results.get("files", [])
    if files:
        return files[0]["id"]
    folder = drive_service.files().create(
        body={"name": name,
              "mimeType": "application/vnd.google-apps.folder",
              "parents": [parent_id]},
        fields="id"
    ).execute()
    log.info(f"Created Drive folder: {name}")
    return folder["id"]
 
 
def upload_to_drive(drive_service, local_path, parent_folder_id):
    file_name = os.path.basename(local_path)
    media     = MediaFileUpload(local_path, resumable=True)
    uploaded  = drive_service.files().create(
        body={"name": file_name, "parents": [parent_folder_id]},
        media_body=media, fields="id"
    ).execute()
    file_id = uploaded["id"]
    drive_service.permissions().create(
        fileId=file_id,
        body={"type": "anyone", "role": "reader"},
    ).execute()
    return f"https://drive.google.com/uc?export=download&id={file_id}"
 
# ---------------------------------------------
#  PLACEHOLDER REPLACEMENT
# ---------------------------------------------
def replace_placeholders(prs, row_data, columns):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    for col in columns:
                        placeholder = f"{{{{{col}}}}}"
                        if placeholder in run.text:
                            value = "" if pd.isna(row_data[col]) else str(row_data[col])
                            run.text = run.text.replace(placeholder, value)
    return prs
 
 
def get_shape_alt_text(shape):
    PML = "http://schemas.openxmlformats.org/presentationml/2006/main"
    for container in (f"{{{PML}}}nvSpPr", f"{{{PML}}}nvPicPr"):
        c = shape._element.find(container)
        if c is not None:
            cNvPr = c.find(f"{{{PML}}}cNvPr")
            if cNvPr is not None:
                return cNvPr.get("descr", "")
    return ""
 
 
def replace_shape_with_qr(slide, qr_path, placeholder_name=QR_PLACEHOLDER):
    for shape in slide.shapes:
        descr = get_shape_alt_text(shape)
        if descr == placeholder_name or shape.name == placeholder_name:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
            shape._element.getparent().remove(shape._element)
            slide.shapes.add_picture(qr_path, left, top, width=width, height=height)
            log.info(f"QR placeholder replaced (matched: '{descr or shape.name}')")
            return True
    log.warning(
        f"QR placeholder '{placeholder_name}' not found. "
        f"Set a shape Alt Text to exactly '{placeholder_name}'."
    )
    return False
 
# ---------------------------------------------
#  PDF / PNG CONVERSION
# ---------------------------------------------
def pptx_to_pdf(powerpoint, pptx_path, pdf_path):
    presentation = powerpoint.Presentations.Open(pptx_path, WithWindow=False)
    po = presentation.PrintOptions
    po.OutputType     = 1
    po.PrintColorType = 1
    po.HighQuality    = True
    presentation.SaveAs(pdf_path, FileFormat=32)
    presentation.Close()
 
 
def pdf_to_png(pdf_path, out_folder, safe_name, dpi=PNG_DPI):
    pages = convert_from_path(pdf_path, dpi=dpi, fmt="png", use_cropbox=True, strict=False)
    png_paths = []
    if len(pages) == 1:
        p = os.path.join(out_folder, f"{safe_name}.png")
        pages[0].save(p, "PNG", optimize=False, compress_level=0)
        png_paths.append(p)
    else:
        for i, page in enumerate(pages, start=1):
            p = os.path.join(out_folder, f"{safe_name}_page{i}.png")
            page.save(p, "PNG", optimize=False, compress_level=0)
            png_paths.append(p)
    return png_paths
 
# ---------------------------------------------
#  EMAIL
# ---------------------------------------------
def load_email_template(file_path, row, columns):
    with open(file_path, "r", encoding="utf-8") as f:
        content = f.read()
    if "\n\n" in content:
        subject_line, body = content.split("\n\n", 1)
    else:
        subject_line = content.split("\n")[0]
        body = ""
    subject = subject_line.replace("Subject:", "").strip()
    for col in columns:
        value = "" if pd.isna(row[col]) else str(row[col])
        subject = subject.replace(f"{{{{{col}}}}}", value)
        body    = body.replace(f"{{{{{col}}}}}", value)
    return subject.strip(), body.strip()
 
 
def send_email(recipient, subject, body, attachment_paths, retries=MAX_EMAIL_RETRIES):
    MIME_MAP = {
        ".pdf":  ("application", "pdf"),
        ".png":  ("image",       "png"),
        ".jpg":  ("image",       "jpeg"),
        ".jpeg": ("image",       "jpeg"),
    }
    for attempt in range(1, retries + 1):
        try:
            msg = EmailMessage()
            msg["From"]    = EMAIL_USER
            msg["To"]      = recipient
            msg["Subject"] = subject
            msg.set_content(body)
            for path in attachment_paths:
                ext = os.path.splitext(path)[1].lower()
                maintype, subtype = MIME_MAP.get(ext, ("application", "octet-stream"))
                with open(path, "rb") as f:
                    msg.add_attachment(f.read(), maintype=maintype, subtype=subtype,
                                       filename=os.path.basename(path))
            with smtplib.SMTP_SSL(ZOHO_SMTP_HOST, ZOHO_SMTP_PORT) as smtp:
                smtp.login(EMAIL_USER, EMAIL_PASS)
                smtp.send_message(msg)
            log.info(f"Email sent to {recipient}")
            return True
        except Exception as e:
            log.warning(f"Email attempt {attempt}/{retries} failed for {recipient}: {e}")
            if attempt < retries:
                time.sleep(RETRY_DELAY_SECONDS)
    log.error(f"All {retries} email attempts failed for {recipient}")
    return False
 
# ---------------------------------------------
#  CERTIFICATE GENERATION (per record)
# ---------------------------------------------
def generate_certificate(idx, row, columns, powerpoint, template_file, out_folder,
                         export_mode, drive_service, drive_subfolder_id):
    name_str  = str(row[NAME_COLUMN]).strip()
    cert_id   = row["CertID"]
    safe_name = re.sub(r'[\\/*?:"<>|]', "_", name_str)
 
    pdf_path = os.path.join(out_folder, f"{safe_name}.pdf")
    png_path = os.path.join(out_folder, f"{safe_name}.png")
 
    # Resume: skip if output already exists
    if (export_mode == "pdf" and os.path.exists(pdf_path)) or \
       (export_mode == "png" and os.path.exists(png_path)):
        log.info(f"Skipped (exists): {name_str}")
        return pdf_path if export_mode == "pdf" else None, [], "", "skipped"
 
    # 1. QR code — white background, minimal border
    qr_url  = f"{VALIDATION_URL}{cert_id}"
    qr_temp = f"temp_qr_{secrets.token_hex(4)}.png"
    qr = qrcode.QRCode(border=1)
    qr.add_data(qr_url)
    qr.make(fit=True)
    qr.make_image(fill_color="black", back_color="white").save(qr_temp)
 
    # 2. Fill PPTX template
    pptx_path = os.path.join(out_folder, f"{safe_name}.pptx")
    prs = Presentation(template_file)
    replace_placeholders(prs, row, columns)
    replace_shape_with_qr(prs.slides[0], qr_temp)
    prs.save(pptx_path)
 
    # 3. PPTX -> PDF
    pptx_to_pdf(powerpoint, pptx_path, pdf_path)
    os.remove(pptx_path)
    os.remove(qr_temp)
 
    # 4. Optional: PDF -> PNG
    final_png_paths = []
    if export_mode == "png":
        final_png_paths = pdf_to_png(pdf_path, out_folder, safe_name)
        os.remove(pdf_path)
        pdf_path = None
 
    # 5. Upload to Google Drive
    upload_path  = pdf_path if export_mode == "pdf" else (final_png_paths[0] if final_png_paths else None)
    download_url = ""
    if upload_path and os.path.exists(upload_path):
        try:
            download_url = upload_to_drive(drive_service, upload_path, drive_subfolder_id)
            log.info(f"Uploaded to Drive: {os.path.basename(upload_path)}")
        except Exception as e:
            log.warning(f"Drive upload failed for {name_str}: {e}")
 
    log.info(f"Generated [{cert_id}]: {name_str}")
    return pdf_path, final_png_paths, download_url, "success"
 
# ---------------------------------------------
#  PROCESS ONE CATEGORY
# ---------------------------------------------
def process_category(sheet_name, template_file, subfolder, export_mode, dry_run,
                     powerpoint, event_title, drive_service, drive_event_folder_id):
    abs_template = os.path.abspath(template_file)
    out_folder   = os.path.join(os.path.abspath(OUTPUT_FOLDER), subfolder)
 
    if not os.path.exists(abs_template):
        log.error(f"Template not found: {abs_template} — skipping '{sheet_name}'")
        return [], [], 0, 0, 0
 
    try:
        df = pd.read_excel(os.path.abspath(DATA_FILE), sheet_name=sheet_name)
    except Exception as e:
        log.error(f"Could not load sheet '{sheet_name}': {e} — skipping")
        return [], [], 0, 0, 0
 
    if NAME_COLUMN not in df.columns:
        log.error(f"'{NAME_COLUMN}' column missing in '{sheet_name}' — skipping")
        return [], [], 0, 0, 0
 
    df = df[df[NAME_COLUMN].notna() & (df[NAME_COLUMN].astype(str).str.strip() != "")]
    if df.empty:
        log.warning(f"Sheet '{sheet_name}' has no valid records — skipping")
        return [], [], 0, 0, 0
 
    if SEND_EMAILS and "Email" not in df.columns:
        log.warning(f"No 'Email' column in '{sheet_name}'. Emails will be skipped.")
    if "Phone" not in df.columns:
        log.warning(f"No 'Phone' column in '{sheet_name}'. phone will be empty in CSV.")
 
    df["CertID"] = [f"{CERT_PREFIX}{generate_unique_suffix()}" for _ in range(len(df))]
 
    os.makedirs(out_folder, exist_ok=True)
 
    # Drive: GDRIVE_FOLDER_ID / event_title / subfolder
    drive_subfolder_id = get_or_create_drive_folder(drive_service, subfolder, drive_event_folder_id)
 
    records = df.iloc[:1] if dry_run else df
    log.info(f"\n[{sheet_name}] -> {subfolder}/ | {len(records)} record(s)")
 
    success_count = skip_count = fail_count = 0
    email_jobs  = []
    master_rows = []
    columns     = list(df.columns)
 
    for idx, row in tqdm(records.iterrows(), total=len(records), desc=f"  {sheet_name}"):
        try:
            pdf_path, png_paths, download_url, status = generate_certificate(
                idx, row, columns, powerpoint, abs_template, out_folder,
                export_mode, drive_service, drive_subfolder_id
            )
 
            if status == "skipped":
                skip_count += 1
            else:
                success_count += 1
                master_rows.append({
                    "recipient_name":   str(row.get(NAME_COLUMN, "")).strip(),
                    "recipient_email":  str(row.get("Email", "")).strip()
                                        if pd.notna(row.get("Email")) else "",
                    "recipient_phone":  str(row.get("Phone", "")).strip()
                                        if pd.notna(row.get("Phone")) else "",
                    "event_title":      event_title,
                    "download_url":     download_url,
                    "certificate_code": str(row.get("CertID", "")).strip(),
                })
 
            attachments = []
            if pdf_path and os.path.exists(pdf_path):
                attachments.append(pdf_path)
            attachments.extend(png_paths)
 
            if (attachments and SEND_EMAILS
                    and "Email" in df.columns
                    and pd.notna(row.get("Email"))
                    and str(row["Email"]).strip()):
                subject, body = load_email_template(EMAIL_TEMPLATE_FILE, row, columns)
                email_jobs.append((str(row["Email"]).strip(), subject, body, attachments))
 
        except Exception as e:
            log.error(f"Failed row {idx} in '{sheet_name}' ({row.get(NAME_COLUMN, '?')}): {e}")
            fail_count += 1
 
    return email_jobs, master_rows, success_count, skip_count, fail_count
 
# ---------------------------------------------
#  MAIN
# ---------------------------------------------
def main():
    dry_run = input("Dry run (first record per sheet only)? (y/n): ").strip().lower() == "y"
 
    print("\nExport format:")
    print("  1 -> PDF")
    print("  2 -> PNG")
    fmt_choice  = input("Choose (1/2) [default: 1]: ").strip()
    export_mode = {"1": "pdf", "2": "png"}.get(fmt_choice, "pdf")
 
    event_title = input("\nEvent title: ").strip()
 
    log.info(f"Export mode : {export_mode.upper()}")
    log.info(f"Event title : {event_title}")
    log.info(f"Starting{'  [DRY RUN]' if dry_run else ''}")
 
    os.makedirs(os.path.abspath(OUTPUT_FOLDER), exist_ok=True)
 
    # Connect to Drive once for the whole run
    log.info("Connecting to Google Drive...")
    drive_service = get_drive_service()
    log.info("Drive connection established.")
 
    # Create: root / event_title /  (subfolders created per category)
    drive_event_folder_id = get_or_create_drive_folder(drive_service, event_title, GDRIVE_FOLDER_ID)
    log.info(f"Drive event folder ready: {event_title}")
 
    all_email_jobs  = []
    all_master_rows = []
    total_success = total_skip = total_fail = 0
 
    powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
    powerpoint.Visible = 1
 
    try:
        for sheet_name, template_file, subfolder in CATEGORIES:
            jobs, rows, s, sk, f = process_category(
                sheet_name, template_file, subfolder, export_mode, dry_run,
                powerpoint, event_title, drive_service, drive_event_folder_id
            )
            all_email_jobs.extend(jobs)
            all_master_rows.extend(rows)
            total_success += s
            total_skip    += sk
            total_fail    += f
    finally:
        powerpoint.Quit()
 
    # CSV — exact column order required
    if all_master_rows:
        csv_name = "certificate_master_log.csv"
        pd.DataFrame(all_master_rows, columns=[
            "recipient_name",
            "recipient_email",
            "recipient_phone",
            "event_title",
            "download_url",
            "certificate_code",
        ]).to_csv(csv_name, index=False)
        log.info(f"Master log saved: {csv_name}")
 
    # Parallel emails
    email_success = email_fail = 0
    if all_email_jobs and SEND_EMAILS:
        log.info(f"Sending {len(all_email_jobs)} email(s) with {EMAIL_THREADS} thread(s)...")
        with ThreadPoolExecutor(max_workers=EMAIL_THREADS) as executor:
            futures = {executor.submit(send_email, *job): job[0] for job in all_email_jobs}
            for future in tqdm(as_completed(futures), total=len(futures), desc="Sending emails"):
                if future.result():
                    email_success += 1
                else:
                    email_fail += 1
 
    drive_event_url = f"https://drive.google.com/drive/folders/{drive_event_folder_id}"
    print("\n" + "-" * 45)
    print("Summary Report")
    print(f"   Generated  : {total_success}")
    print(f"   Skipped    : {total_skip}  (already existed)")
    print(f"   Failed     : {total_fail}")
    if all_email_jobs:
        print(f"   Emails     : {email_success} sent / {email_fail} failed")
    print(f"\nLocal output:")
    for _, _, subfolder in CATEGORIES:
        print(f"   {OUTPUT_FOLDER}/{subfolder}/")
    print(f"\nDrive structure:")
    print(f"   (root folder)/")
    print(f"   └── {event_title}/")
    for _, _, subfolder in CATEGORIES:
        print(f"         ├── {subfolder}/")
    print(f"\nDrive link : {drive_event_url}")
    print(f"CSV        : certificate_master_log.csv")
    print("-" * 45)
 
 
if __name__ == "__main__":
    main()